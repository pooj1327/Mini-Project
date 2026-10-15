"""
resume/resume_parser.py

Extracts text from PDF and DOCX resumes.
"""

from pathlib import Path

import fitz  # PyMuPDF
from docx import Document


class ResumeParser:

    def __init__(self, file_path: str):
        self.file_path = Path(file_path)

    # ---------------------------------------
    # Read PDF
    # ---------------------------------------

    def _read_pdf(self):

        text = ""

        pdf = fitz.open(self.file_path)

        for page in pdf:

            text += page.get_text()

        pdf.close()

        return text

    # ---------------------------------------
    # Read DOCX
    # ---------------------------------------

    def _read_docx(self):

        doc = Document(self.file_path)

        text = "\n".join(
            paragraph.text
            for paragraph in doc.paragraphs
        )

        return text

    # ---------------------------------------
    # Extract Text
    # ---------------------------------------

    def extract_text(self):

        suffix = self.file_path.suffix.lower()

        if suffix == ".pdf":
            return self._read_pdf()

        elif suffix == ".docx":
            return self._read_docx()

        raise ValueError(
            "Unsupported Resume Format."
        )

    # ---------------------------------------
    # Clean Text
    # ---------------------------------------

    @staticmethod
    def clean_text(text):

        text = text.replace("\t", " ")

        text = text.replace("\r", " ")

        text = text.replace("\n\n", "\n")

        return text.strip()

    # ---------------------------------------
    # Guess Name by Font Size
    # ---------------------------------------
    # The name is almost always the largest text on the page in
    # real resume templates - this is far more reliable than
    # assuming it's simply the first line, especially for PDFs
    # with multi-column layouts or graphic headers where the text
    # extraction order doesn't match visual reading order.

    def guess_name(self):

        suffix = self.file_path.suffix.lower()

        try:

            if suffix == ".pdf":
                return self._guess_name_pdf()

            elif suffix == ".docx":
                return self._guess_name_docx()

        except Exception:
            return None

        return None

    def _guess_name_pdf(self):

        pdf = fitz.open(self.file_path)

        page = pdf[0]

        blocks = page.get_text("dict")["blocks"]

        candidates = []

        for block in blocks:

            for line in block.get("lines", []):

                for span in line.get("spans", []):

                    text = span["text"].strip()

                    if not text or "@" in text:
                        continue

                    if len(text) < 3 or len(text) > 40:
                        continue

                    candidates.append((span["size"], text))

        pdf.close()

        if not candidates:
            return None

        candidates.sort(key=lambda c: c[0], reverse=True)

        return candidates[0][1]

    def _guess_name_docx(self):

        doc = Document(self.file_path)

        candidates = []

        for paragraph in doc.paragraphs[:15]:

            for run in paragraph.runs:

                text = run.text.strip()

                if not text or "@" in text:
                    continue

                if len(text) < 3 or len(text) > 40:
                    continue

                size = run.font.size.pt if run.font.size else 11

                candidates.append((size, text))

        if not candidates:
            return None

        candidates.sort(key=lambda c: c[0], reverse=True)

        return candidates[0][1]

    # ---------------------------------------
    # Parse Resume
    # ---------------------------------------

    def parse(self):

        text = self.extract_text()

        text = self.clean_text(text)

        return text